#!/usr/bin/env python3
"""
Regenerate augmentation_strategy.docx from augmentation_strategy.md so the two
never drift apart. Handles the Markdown used there: headings, paragraphs,
bullet/numbered lists, pipe tables, fenced code, **bold** and `code`.

Requires python-docx (pip install python-docx).
Usage: python build_docx.py
"""
from pathlib import Path
import re
import docx
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_runs(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(part)


def cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main():
    lines = (HERE / "augmentation_strategy.md").read_text().splitlines()
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10)
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            p = d.add_paragraph()
            r = p.add_run("\n".join(lines[i + 1:j]))
            r.font.name, r.font.size = "Consolas", Pt(9)
            i = j + 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            rows = [cells(b) for b in block if not re.match(r"^\|\s*-", b)]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    par = t.cell(r, c).paragraphs[0]
                    add_runs(par, text)
                    if r == 0:
                        for run in par.runs:
                            run.bold = True
            continue
        if line.startswith("# "):
            d.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            d.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            d.add_heading(line[4:], level=2)
        elif re.match(r"^\s*- ", line):
            add_runs(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*- ", "", line))
        elif re.match(r"^\d+\. ", line):
            add_runs(d.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.startswith("  ") and d.paragraphs and line.strip():
            add_runs(d.paragraphs[-1], " " + line.strip())  # continuation of a list item
        elif line.strip():
            para = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||-|\d+\. |```|\s+- )", lines[i + 1]):
                i += 1; para.append(lines[i])
            add_runs(d.add_paragraph(), " ".join(p.strip() for p in para))
        i += 1
    d.save(HERE / "augmentation_strategy.docx")
    print("wrote augmentation_strategy.docx")


if __name__ == "__main__":
    main()
